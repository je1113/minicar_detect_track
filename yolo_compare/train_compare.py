#!/usr/bin/env python3
"""Roboflow AMR 데이터셋으로 YOLOv8 ~ YOLO26을 같은 조건에서 학습하고 결과를 엑셀로 정리한다.

사용법:
    export ROBOFLOW_API_KEY="..."
    python3 train_compare.py                       # 전체 모델
    python3 train_compare.py --models yolo26n.pt   # 일부만
    python3 train_compare.py --epochs 50 --batch 8
"""
import argparse
import gc
import os
import time
from pathlib import Path

import torch
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from roboflow import Roboflow
from ultralytics import YOLO
from ultralytics.utils.torch_utils import get_flops, get_num_params

ROOT = Path(__file__).resolve().parent

DEFAULT_MODELS = [
    "yolov8n.pt",
    "yolov9t.pt",
    "yolov10n.pt",
    "yolo11n.pt",
    "yolo12n.pt",
    "yolo26n.pt",
]

COLUMNS = [
    "Model", "Params(M)", "GFLOPs", "Precision", "Recall",
    "mAP50", "mAP50-95", "Inference(ms/img)", "Train time(min)", "Error",
]


def download_dataset():
    data_yaml = ROOT / "datasets" / "amr-v1" / "data.yaml"
    if data_yaml.exists():
        return data_yaml
    api_key = os.environ.get("ROBOFLOW_API_KEY")
    if not api_key:
        raise SystemExit("ROBOFLOW_API_KEY 환경변수를 먼저 설정하세요.")
    rf = Roboflow(api_key=api_key)
    project = rf.workspace("-jg3fi").project("amr-5qizb")
    dataset = project.version(1).download("yolov8", location=str(ROOT / "datasets" / "amr-v1"))
    return Path(dataset.location) / "data.yaml"


def train_and_eval(weights, data_yaml, args):
    name = Path(weights).stem
    model = YOLO(weights)
    t0 = time.time()
    model.train(
        data=str(data_yaml), epochs=args.epochs, imgsz=args.imgsz, batch=args.batch,
        seed=0, deterministic=True, project=str(ROOT / "runs"), name=name, exist_ok=True,
    )
    train_min = (time.time() - t0) / 60

    best = YOLO(ROOT / "runs" / name / "weights" / "best.pt")
    m = best.val(data=str(data_yaml), imgsz=args.imgsz, batch=args.batch, split="val",
                 project=str(ROOT / "runs"), name=f"{name}_val", exist_ok=True)
    params = get_num_params(best.model)
    gflops = get_flops(best.model, args.imgsz)

    return {
        "Model": name,
        "Params(M)": round(params / 1e6, 2),
        "GFLOPs": round(gflops, 1),
        "Precision": round(float(m.box.mp), 4),
        "Recall": round(float(m.box.mr), 4),
        "mAP50": round(float(m.box.map50), 4),
        "mAP50-95": round(float(m.box.map), 4),
        "Inference(ms/img)": round(m.speed["inference"], 2),
        "Train time(min)": round(train_min, 1),
    }


def open_workbook(path, sheets):
    """기존 엑셀을 열고 다시 쓸 시트만 지운다. 다른 스크립트가 만든 시트는 그대로 둔다."""
    if Path(path).exists():
        wb = load_workbook(path)
    else:
        wb = Workbook()
        wb.remove(wb.active)
    for name in sheets:
        if name in wb.sheetnames:
            del wb[name]
    return wb


def autofit(ws):
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(width + 2, 60)
    ws.freeze_panes = "A2"


def save_excel(rows, path, args):
    wb = open_workbook(path, ["comparison", "settings"])
    ws = wb.create_sheet("comparison", 0)
    ws.append(COLUMNS)
    for row in rows:
        ws.append([row.get(c) for c in COLUMNS])

    # mAP50-95 최고 모델 강조
    ok = [r for r in rows if r.get("mAP50-95") is not None]
    if ok:
        best_name = max(ok, key=lambda r: r["mAP50-95"])["Model"]
        for r in ws.iter_rows(min_row=2):
            if r[0].value == best_name:
                for cell in r:
                    cell.fill = PatternFill("solid", fgColor="C6EFCE")

    autofit(ws)

    info = wb.create_sheet("settings", 1)
    for k, v in [("epochs", args.epochs), ("imgsz", args.imgsz), ("batch", args.batch),
                 ("seed", 0), ("dataset", "-jg3fi/amr-5qizb v1")]:
        info.append([k, v])

    wb.save(path)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--models", nargs="+", default=DEFAULT_MODELS)
    parser.add_argument("--epochs", type=int, default=100)
    parser.add_argument("--imgsz", type=int, default=640)
    parser.add_argument("--batch", type=int, default=16)
    parser.add_argument("--out", default=str(ROOT / "yolo_comparison.xlsx"))
    args = parser.parse_args()

    data_yaml = download_dataset()
    rows = []
    for weights in args.models:
        print(f"\n===== {weights} =====")
        try:
            rows.append(train_and_eval(weights, data_yaml, args))
        except Exception as e:
            print(f"[{weights}] 실패: {e}")
            rows.append({"Model": Path(weights).stem, "Error": str(e)})
        finally:
            gc.collect()
            if torch.cuda.is_available():
                torch.cuda.empty_cache()
        # 중간에 끊겨도 결과가 남도록 모델마다 저장
        save_excel(rows, args.out, args)

    print(f"\n결과 저장: {args.out}")
    for r in rows:
        print(r)


if __name__ == "__main__":
    main()

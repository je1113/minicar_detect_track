#!/usr/bin/env python3
"""YOLO 모델 크기 × optimizer 조합을 여러 seed로 학습하고 yolo_comparison.xlsx에 시트로 추가한다.

사용법:
    export ROBOFLOW_API_KEY="..."
    python3 train_sweep.py                                   # yolo26 n/s/m × SGD/AdamW/MuSGD × seed 0,1,2
    python3 train_sweep.py --sizes n s --optimizers SGD AdamW --seeds 0
    python3 train_sweep.py --family yolo11 --lr0 0.005       # 모든 optimizer에 같은 lr0 적용

중간에 끊겨도 다시 실행하면 runs/sweep/results.jsonl에 기록된 조합은 건너뛴다.
"""
import argparse
import gc
import json
import statistics
import time
from pathlib import Path

import torch
from openpyxl.styles import PatternFill
from ultralytics import YOLO
from ultralytics.utils.torch_utils import get_flops, get_num_params

from train_compare import autofit, download_dataset, open_workbook

ROOT = Path(__file__).resolve().parent
SWEEP_DIR = ROOT / "runs" / "sweep"
LOG = SWEEP_DIR / "results.jsonl"

# optimizer를 직접 지정하면 ultralytics는 lr0(기본 0.01)를 그대로 쓴다.
# Adam 계열에 0.01은 너무 커서 optimizer별 통상 값을 따로 둔다. "auto"는 ultralytics가 직접 고른다.
DEFAULT_LR0 = {"SGD": 0.01, "MuSGD": 0.01, "AdamW": 0.001, "Adam": 0.001, "NAdam": 0.001,
               "RAdam": 0.001, "Adamax": 0.002, "RMSprop": 0.0001}

METRICS = ["Val P", "Val R", "Val mAP50", "Val mAP50-95", "Test mAP50", "Test mAP50-95",
           "Inference(ms/img)", "Train time(min)"]
RUN_COLUMNS = ["Model", "Size", "Optimizer", "lr0", "Epochs", "imgsz", "Seed", "Batch", "Params(M)", "GFLOPs",
               *METRICS, "Error"]
SUMMARY_KEYS = ["Model", "Size", "Optimizer", "lr0", "Epochs", "imgsz", "Params(M)", "GFLOPs"]


def run_key(r):
    return (r["Model"], r["Optimizer"], r["lr0"], r["Epochs"], r["imgsz"], r["Seed"])


def load_log():
    if not LOG.exists():
        return []
    return [json.loads(line) for line in LOG.read_text().splitlines() if line.strip()]


def is_oom(e):
    return isinstance(e, torch.cuda.OutOfMemoryError) or "out of memory" in str(e).lower()


def free_gpu():
    gc.collect()
    if torch.cuda.is_available():
        torch.cuda.empty_cache()


def train_one(weights, optimizer, lr0, seed, batch, data_yaml, args):
    name = f"{Path(weights).stem}_{optimizer}_lr{lr0}_e{args.epochs}_i{args.imgsz}_s{seed}"
    model = YOLO(weights)
    train_args = dict(data=str(data_yaml), epochs=args.epochs, imgsz=args.imgsz, batch=batch,
                      optimizer=optimizer, seed=seed, deterministic=True,
                      project=str(SWEEP_DIR), name=name, exist_ok=True)
    if optimizer != "auto":
        train_args["lr0"] = lr0
    t0 = time.time()
    model.train(**train_args)
    train_min = (time.time() - t0) / 60

    best = YOLO(SWEEP_DIR / name / "weights" / "best.pt")
    val_kw = dict(data=str(data_yaml), imgsz=args.imgsz, batch=batch,
                  project=str(SWEEP_DIR), exist_ok=True)
    v = best.val(split="val", name=f"{name}_val", **val_kw)
    t = best.val(split="test", name=f"{name}_test", **val_kw)

    return {
        "Params(M)": round(get_num_params(best.model) / 1e6, 2),
        "GFLOPs": round(get_flops(best.model, args.imgsz), 1),
        "Val P": round(float(v.box.mp), 4),
        "Val R": round(float(v.box.mr), 4),
        "Val mAP50": round(float(v.box.map50), 4),
        "Val mAP50-95": round(float(v.box.map), 4),
        "Test mAP50": round(float(t.box.map50), 4),
        "Test mAP50-95": round(float(t.box.map), 4),
        "Inference(ms/img)": round(v.speed["inference"], 2),
        "Train time(min)": round(train_min, 1),
    }


def run_with_oom_retry(weights, optimizer, lr0, seed, data_yaml, args):
    """GPU 메모리가 부족하면 batch를 절반으로 줄여 다시 시도한다. 실제 사용한 batch를 함께 반환한다."""
    batch = args.batch
    while True:
        try:
            return batch, train_one(weights, optimizer, lr0, seed, batch, data_yaml, args)
        except Exception as e:
            if not is_oom(e) or batch <= 2:
                raise
            free_gpu()
            batch //= 2
            print(f"[{weights}] GPU 메모리 부족 → batch {batch}로 재시도")


def summarize(rows):
    """같은 (모델, optimizer, lr0, epochs, imgsz) 조합의 seed별 결과를 평균 ± 표준편차로 묶는다."""
    groups = {}
    for r in rows:
        if r.get("Error"):
            continue
        groups.setdefault(run_key(r)[:-1], []).append(r)

    out = []
    for runs in groups.values():
        s = {k: runs[0][k] for k in SUMMARY_KEYS}
        s["Seeds"] = len(runs)
        for m in METRICS:
            vals = [r[m] for r in runs]
            s[f"{m} mean"] = round(statistics.mean(vals), 4)
            s[f"{m} std"] = round(statistics.stdev(vals), 4) if len(vals) > 1 else 0.0
        out.append(s)
    return sorted(out, key=lambda s: s["Val mAP50-95 mean"], reverse=True)


def save_excel(rows, path, args):
    wb = open_workbook(path, ["sweep_summary", "sweep_runs", "sweep_settings"])
    green = PatternFill("solid", fgColor="C6EFCE")

    summary = summarize(rows)
    ws = wb.create_sheet("sweep_summary")
    cols = SUMMARY_KEYS + ["Seeds"] + [f"{m} {s}" for m in METRICS for s in ("mean", "std")]
    ws.append(cols)
    for s in summary:
        ws.append([s.get(c) for c in cols])
    if summary:  # Val mAP50-95 평균 1위 강조
        for cell in ws[2]:
            cell.fill = green
    autofit(ws)

    ws = wb.create_sheet("sweep_runs")
    ws.append(RUN_COLUMNS)
    for r in rows:
        ws.append([r.get(c) for c in RUN_COLUMNS])
    autofit(ws)

    info = wb.create_sheet("sweep_settings")
    for k, v in [("family", args.family), ("sizes", " ".join(args.sizes)),
                 ("optimizers", " ".join(args.optimizers)), ("seeds", " ".join(map(str, args.seeds))),
                 ("lr0", args.lr0 if args.lr0 is not None else "optimizer별 기본값"),
                 ("epochs", args.epochs), ("imgsz", args.imgsz), ("batch", args.batch),
                 ("dataset", "-jg3fi/amr-5qizb v1")]:
        info.append([k, v])

    wb.save(path)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--family", default="yolo26", help="가중치 이름 앞부분 (yolo26, yolo11, yolov8 ...)")
    parser.add_argument("--sizes", nargs="+", default=["n", "s", "m"])
    parser.add_argument("--optimizers", nargs="+", default=["SGD", "AdamW", "MuSGD"])
    parser.add_argument("--seeds", nargs="+", type=int, default=[0, 1, 2])
    parser.add_argument("--lr0", type=float, default=None, help="지정하면 모든 optimizer에 이 값 사용")
    parser.add_argument("--epochs", type=int, default=100)
    parser.add_argument("--imgsz", type=int, default=640)
    parser.add_argument("--batch", type=int, default=16)
    parser.add_argument("--out", default=str(ROOT / "yolo_comparison.xlsx"))
    args = parser.parse_args()

    data_yaml = download_dataset()
    SWEEP_DIR.mkdir(parents=True, exist_ok=True)
    rows = load_log()
    done = {run_key(r) for r in rows if not r.get("Error")}
    rows = [r for r in rows if not r.get("Error")]  # 실패했던 조합은 다시 시도

    plan = [(size, opt, seed) for size in args.sizes for opt in args.optimizers for seed in args.seeds]
    for i, (size, opt, seed) in enumerate(plan, 1):
        weights = f"{args.family}{size}.pt"
        lr0 = "auto" if opt == "auto" else (args.lr0 if args.lr0 is not None else DEFAULT_LR0.get(opt, 0.01))
        row = {"Model": Path(weights).stem, "Size": size, "Optimizer": opt, "lr0": lr0,
               "Epochs": args.epochs, "imgsz": args.imgsz, "Seed": seed}
        if run_key(row) in done:
            print(f"[{i}/{len(plan)}] {weights} {opt} lr0={lr0} seed={seed} 이미 완료 → 건너뜀")
            continue

        print(f"\n===== [{i}/{len(plan)}] {weights} | {opt} lr0={lr0} | seed {seed} =====")
        try:
            batch, metrics = run_with_oom_retry(weights, opt, lr0, seed, data_yaml, args)
            row.update(Batch=batch, **metrics)
        except Exception as e:
            print(f"[{weights} {opt} seed{seed}] 실패: {e}")
            row.update(Batch=args.batch, Error=str(e))
        finally:
            free_gpu()

        rows.append(row)
        with LOG.open("a") as f:
            f.write(json.dumps(row, ensure_ascii=False) + "\n")
        save_excel(rows, args.out, args)

    save_excel(rows, args.out, args)
    print(f"\n결과 저장: {args.out} (sweep_summary / sweep_runs 시트)")
    for s in summarize(rows):
        print(f"{s['Model']:10s} {s['Optimizer']:6s} lr0={s['lr0']}: "
              f"val mAP50-95 {s['Val mAP50-95 mean']:.4f} ± {s['Val mAP50-95 std']:.4f}, "
              f"test {s['Test mAP50-95 mean']:.4f} ± {s['Test mAP50-95 std']:.4f}")


if __name__ == "__main__":
    main()
